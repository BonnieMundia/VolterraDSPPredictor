"""Builds the VolterraDSPPredictor presentation in two formats from one slide spec:

    VolterraDSPPredictor_Slides.pptx   editable deck with speaker notes (python-pptx)
    VolterraDSPPredictor_Slides.html   same layout, printed to PDF with headless Chrome

Usage:  python3 -m venv /tmp/docvenv && /tmp/docvenv/bin/pip install python-pptx
        /tmp/docvenv/bin/python make_slides.py
"""
import html
import os
import struct

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = 13.333, 7.5                      # 16:9 slide in inches
FONT = "Helvetica Neue"
INK, MUTED, ACCENT = "1B1F24", "5A6370", "1F5FBF"
BLUE, ORANGE, GREEN = "3D7BE0", "E0662F", "3BA272"
TINT, ORANGE_TINT, GREEN_TINT = "EEF3FB", "FCEFE8", "EAF6EF"


def png_size(path):
    with open(path, "rb") as f:
        f.read(16)
        return struct.unpack(">II", f.read(8))


# ---------------------------------------------------------------- slide spec
# Elements (all coordinates in inches):
#   ("text", x, y, w, h, [paragraph, ...], opts)  paragraph = str | (str, {size, bold, color, bullet})
#   ("image", x, y, w, path, caption)
#   ("tile", x, y, w, h, big, label, color, fill)
#   ("box", x, y, w, h, title, subtitle, color, fill)
#   ("arrow", x1, y1, x2, y2, label)

def bullets(*items, size=20):
    return [(t, {"size": size, "bullet": True}) for t in items]


SLIDES = [
    {   # 1
        "title": None,
        "elements": [
            ("text", 0.9, 1.75, 11.5, 1.2, [("VolterraDSPPredictor", {"size": 54, "bold": True, "color": INK})], {}),
            ("text", 0.9, 2.95, 11.5, 1.2, [("Teaching a radio amplifier to stay clean, using software that learns", {"size": 26, "color": ACCENT})], {}),
            ("text", 0.9, 4.3, 11.5, 1.0, [("A real-time simulator for adaptive digital predistortion on Apple Silicon", {"size": 18, "color": MUTED})], {}),
            ("text", 0.9, 6.2, 11.5, 0.6, [("Boniface Mundia  ·  October 2026  ·  github.com/BonnieMundia/VolterraDSPPredictor", {"size": 14, "color": MUTED})], {}),
        ],
        "notes": "Introduce yourself. One-sentence summary: this project shows, live, how software can clean up the "
                 "distortion that radio amplifiers create, so they can run more efficiently.",
    },
    {   # 2
        "title": "The problem: amplifiers distort when pushed hard",
        "elements": [
            ("text", 0.7, 1.55, 6.3, 5.0, bullets(
                "Every phone mast, Wi-Fi router and phone has a power amplifier that boosts the signal before the antenna",
                "Like a speaker turned up too loud, it distorts near its limit: the loudest moments get squashed and twisted",
                "It also has memory: what it does now depends on what it just did",
                "Result: errors in our own signal and noise that spills into other people's channels",
                size=19), {}),
            ("image", 7.3, 1.65, 5.4, "slides_assets/amam.png",
             "Measured output (orange) falls below the ideal straight line as the input gets louder"),
        ],
        "notes": "Use the speaker analogy. On the chart, a perfect amplifier would be a straight line. The orange dots "
                 "bend down at high input: that is compression. The spread of dots is the memory effect.",
    },
    {   # 3
        "title": "Why it matters",
        "elements": [
            ("tile", 0.7, 1.8, 3.8, 3.1, "Energy", "The amplifier is the biggest power user in a radio. It is most efficient near its limit, which is exactly where it distorts.", ACCENT, TINT),
            ("tile", 4.77, 1.8, 3.8, 3.1, "Neighbours", "Distortion leaks into neighbouring frequency channels. Regulators set strict limits on this leakage.", ORANGE, ORANGE_TINT),
            ("tile", 8.83, 1.8, 3.8, 3.1, "Quality", "A distorted signal causes bit errors, so users get slower, less reliable data.", GREEN, GREEN_TINT),
            ("text", 0.7, 5.35, 11.9, 1.2, [("Today's common fix is to turn the volume down (\"back-off\"). It works, but it wastes energy and money on every transmitter.", {"size": 20, "color": INK})], {}),
        ],
        "notes": "Three reasons: energy bills and battery life, legal limits on interference, and user experience. "
                 "Back-off is like never driving above second gear: safe but wasteful.",
    },
    {   # 4
        "title": "The idea: pre-distort the signal on purpose",
        "elements": [
            ("box", 0.6, 2.0, 2.4, 1.1, "Clean signal", "what we want to send", MUTED, "F4F5F7"),
            ("box", 3.7, 2.0, 2.6, 1.1, "Predistorter (DPD)", "bends it the opposite way", ACCENT, TINT),
            ("box", 7.0, 2.0, 2.4, 1.1, "Amplifier", "bends it back", ORANGE, ORANGE_TINT),
            ("box", 10.1, 2.0, 2.6, 1.1, "Clean, strong output", "distortions cancel", GREEN, GREEN_TINT),
            ("arrow", 3.0, 2.55, 3.7, 2.55, ""),
            ("arrow", 6.3, 2.55, 7.0, 2.55, ""),
            ("arrow", 9.4, 2.55, 10.1, 2.55, ""),
            ("box", 5.6, 3.75, 3.3, 1.0, "Measure & learn", "compare output with the ideal", ACCENT, "FFFFFF"),
            ("arrow", 11.4, 3.1, 8.9, 4.25, ""),
            ("arrow", 5.6, 4.25, 5.0, 3.1, ""),
            ("text", 0.7, 5.15, 11.9, 1.8, bullets(
                "Like glasses: the lens bends light the opposite way your eye does, so you end up seeing clearly",
                "The amplifier changes with heat and age, so the predistorter keeps learning while it works",
                size=20), {}),
        ],
        "notes": "This is the core idea of digital predistortion. The feedback loop is what makes it adaptive: we "
                 "listen to the amplifier's output and keep adjusting the correction.",
    },
    {   # 5
        "title": "The gap this project addresses",
        "elements": [
            ("text", 0.7, 1.55, 11.9, 5.4, bullets(
                "Most research is done offline: record data, compute, plot. You can't watch the learning happen or change things live.",
                "The most accurate model (the full Volterra series) is often dismissed as too slow. Few tools measure what it really costs.",
                "Modern chips where the CPU and GPU share memory, like Apple Silicon, are rarely studied for this job.",
                "Learning methods are compared in different papers with different setups, so the comparison isn't fair.",
                "Few open, reproducible tools exist for teaching: commercial toolboxes are closed and lab hardware is expensive.",
                size=21), {}),
        ],
        "notes": "These five gaps are the motivation. Each one maps to a feature of the app: live loop, full Volterra "
                 "model, CPU+GPU, same setup for every algorithm, and open source with tests.",
    },
    {   # 6
        "title": "What we built",
        "elements": [
            ("image", 0.6, 1.45, 6.7, "screenshots/live_default.png", "The Live Dashboard updates ten times a second"),
            ("text", 7.7, 1.5, 5.1, 5.4, bullets(
                "A Mac app that simulates the whole loop live: signal, predistorter, amplifier, learning, measurement",
                "Change any setting while it runs and see the effect instantly",
                "Four tabs: signal & amplifier, model engine, learning lab, live dashboard",
                "Written in Swift on Apple's Accelerate and Metal, with no external libraries and a 66-case test suite",
                size=17), {}),
        ],
        "notes": "Give a quick tour of the screenshot: tiles at the top, spectrum in the middle, error trace and "
                 "constellation below, per-stage timing at the bottom. If possible, show the app live here.",
    },
    {   # 7
        "title": "How it learns",
        "elements": [
            ("text", 0.7, 1.5, 5.6, 5.4, [
                ("A model with memory", {"size": 20, "bold": True, "color": ACCENT}),
                ("The Volterra model looks at the last few samples and how they combine, not just the current one.", {"size": 17}),
                ("Three ways to learn its settings", {"size": 20, "bold": True, "color": ACCENT}),
                ("NLMS: small, cheap steps. Simple, but slow to settle.", {"size": 17, "bullet": True}),
                ("RLS: smart steps. Fast to settle, but heavy maths on every sample.", {"size": 17, "bullet": True}),
                ("Block RLS: the same smart answer, computed in batches. Fast and accurate, and the default.", {"size": 17, "bullet": True}),
            ], {}),
            ("image", 6.6, 1.6, 6.1, "slides_assets/convergence.png",
             "Switching the predistorter on: the error (blue) drops about 15 dB below the uncorrected amplifier (orange) and stays there"),
        ],
        "notes": "Memory depth M is how many past samples the model looks at. The chart shows the moment adaptation "
                 "starts: the blue line falls and settles within a few frames of signal.",
    },
    {   # 8
        "title": "Results: a much cleaner signal",
        "elements": [
            ("tile", 0.7, 1.5, 3.4, 1.7, "23× less", "leakage to neighbours\n−31.7 → −45.3 dBc", ACCENT, TINT),
            ("tile", 0.7, 3.35, 3.4, 1.7, "40× less", "error in our own signal\n−23.4 → −39.4 dB", GREEN, GREEN_TINT),
            ("tile", 0.7, 5.2, 3.4, 1.7, "7× lower", "symbol error (EVM)\n4.9 % → 0.7 %", ORANGE, ORANGE_TINT),
            ("image", 4.4, 1.5, 8.3, "slides_assets/spectrum.png",
             "Spectrum: orange = amplifier alone, blue = with predistortion, green = ideal"),
            ("image", 4.4, 4.8, 3.6, "slides_assets/ampm_before_after.png", ""),
            ("text", 8.3, 4.85, 4.4, 2.0, [("The twist is gone too: with predistortion the amplifier's phase error (blue) stays flat at 0°. Without it (orange) it wanders by 10–15°.", {"size": 16, "color": INK})], {}),
        ],
        "notes": "Point at the shoulders either side of the main channel: they drop sharply. Default scenario: 16-QAM, 25 MHz wide, Rapp amplifier with memory, 44 model coefficients. "
                 "23× is 13.6 dB and 40× is 16 dB.",
    },
    {   # 9
        "title": "Making it fast",
        "elements": [
            ("tile", 0.7, 1.6, 3.8, 2.7, "23× faster", "Block RLS versus classic RLS (M = 8), with exactly the same result", ACCENT, TINT),
            ("tile", 4.77, 1.6, 3.8, 2.7, "26× faster", "Keep only the important model terms and learn on 1 frame in 8. Same leakage result.", GREEN, GREEN_TINT),
            ("tile", 8.83, 1.6, 3.8, 2.7, "Auto CPU / GPU", "The app times both processors while it runs and picks the faster one. The GPU wins for big models.", ORANGE, ORANGE_TINT),
            ("text", 0.7, 4.75, 11.9, 2.0, bullets(
                "All numbers measured on an Apple M5 laptop chip (Release build, 4,096-sample frames)",
                "Lesson: the \"too slow\" full model becomes practical with smart maths, not just faster hardware",
                size=20), {}),
        ],
        "notes": "Classic RLS needs about 349 ms per frame at memory depth 8; Block RLS needs 15 ms. Pruning plus "
                 "adapting every 8th frame brings it to 0.58 ms. The GPU takes over at memory depth 12 and above.",
    },
    {   # 10
        "title": "Limits, next steps, and try it yourself",
        "elements": [
            ("text", 0.7, 1.5, 5.8, 4.6, [
                ("Honest limits", {"size": 21, "bold": True, "color": ORANGE}),
                ("Simulated amplifier, not real hardware yet", {"size": 18, "bullet": True}),
                ("Third-order model only (no 5th/7th order)", {"size": 18, "bullet": True}),
                ("Not yet at full radio speed (122.88 MS/s)", {"size": 18, "bullet": True}),
                ("Next steps", {"size": 21, "bold": True, "color": ACCENT}),
                ("Test on a real amplifier with a software radio", {"size": 18, "bullet": True}),
                ("Higher-order and neural-network models", {"size": 18, "bullet": True}),
            ], {}),
            ("tile", 7.0, 1.6, 5.6, 3.0, "Try it", "github.com/BonnieMundia/VolterraDSPPredictor\nOpen in Xcode 27 on an Apple Silicon Mac, then press ⌘R", GREEN, GREEN_TINT),
            ("text", 7.0, 5.0, 5.6, 1.2, [("Thank you. Questions?", {"size": 30, "bold": True, "color": INK})], {"align": "center"}),
        ],
        "notes": "Be upfront about limits; they are also the roadmap. Invite people to clone the repo and try the "
                 "scenarios from the documentation.",
    },
]


# ---------------------------------------------------------------- PPTX renderer

def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, size, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def set_bullet(paragraph, on):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if on:
        pPr.set("marL", str(Inches(0.32)))
        pPr.set("indent", str(-Inches(0.26)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    else:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def add_text(slide, x, y, w, h, paragraphs, align=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.04))
    for i, p in enumerate(paragraphs):
        text, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_bullet(para, o.get("bullet", False))
        para.space_after = Pt(o.get("size", 20) * 0.55)
        if align == "center":
            para.alignment = PP_ALIGN.CENTER
        lines = text.split("\n")
        for j, line in enumerate(lines):
            run = para.add_run()
            run.text = line
            style_run(run, o.get("size", 20), o.get("bold", False), o.get("color", INK))
            if j < len(lines) - 1:
                para.add_line_break()
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.shadow.inherit = False
    return s


def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.author = "Boniface Mundia"
    prs.core_properties.title = "VolterraDSPPredictor"
    blank = prs.slide_layouts[6]
    for n, spec in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(blank)
        if spec["title"]:
            add_text(s, 0.7, 0.42, 11.9, 0.8, [(spec["title"], {"size": 32, "bold": True})])
            rect(s, 0.72, 1.18, 1.2, 0.06, ACCENT, shape=MSO_SHAPE.RECTANGLE)
            add_text(s, 0.7, 7.0, 6, 0.35, [("VolterraDSPPredictor", {"size": 11, "color": MUTED})])
            add_text(s, 11.6, 7.0, 1.0, 0.35, [(str(n), {"size": 11, "color": MUTED})], align="center")
        else:
            rect(s, 0, 0, 0.35, H, ACCENT, shape=MSO_SHAPE.RECTANGLE)
        for el in spec["elements"]:
            kind = el[0]
            if kind == "text":
                _, x, y, w, h, paras, opts = el
                add_text(s, x, y, w, h, paras, align=opts.get("align"))
            elif kind == "image":
                _, x, y, w, img, caption = el
                pw, ph = png_size(os.path.join(HERE, img))
                h = w * ph / pw
                s.shapes.add_picture(os.path.join(HERE, img), Inches(x), Inches(y), width=Inches(w))
                add_text(s, x, y + h + 0.08, w, 0.7, [(caption, {"size": 13, "color": MUTED})])
            elif kind == "tile":
                _, x, y, w, h, big, label, color, fill = el
                rect(s, x, y, w, h, fill)
                rect(s, x, y, 0.09, h, color, shape=MSO_SHAPE.RECTANGLE)
                add_text(s, x + 0.3, y + 0.22, w - 0.5, h - 0.3, [
                    (big, {"size": 30 if len(big) < 14 else 26, "bold": True, "color": color}),
                    (label, {"size": 15, "color": INK}),
                ])
            elif kind == "box":
                _, x, y, w, h, title, sub, color, fill = el
                rect(s, x, y, w, h, fill, line=color)
                add_text(s, x, y, w, h, [(title, {"size": 17, "bold": True, "color": INK}),
                                         (sub, {"size": 12, "color": MUTED})],
                         align="center", anchor=MSO_ANCHOR.MIDDLE)
            elif kind == "arrow":
                _, x1, y1, x2, y2, _label = el
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
                c.line.color.rgb = rgb(INK)
                c.line.width = Pt(2)
                ln = c.line._get_or_add_ln()
                ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        s.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(path)


# ---------------------------------------------------------------- HTML renderer (for PDF)

PX = 96  # px per inch


def css_box(x, y, w, h=None):
    s = f"left:{x*PX:.1f}px;top:{y*PX:.1f}px;width:{w*PX:.1f}px;"
    return s + (f"height:{h*PX:.1f}px;" if h is not None else "")


def html_paras(paras, align=None):
    out = []
    for p in paras:
        text, o = (p, {}) if isinstance(p, str) else p
        size = o.get("size", 20) * PX / 72
        style = f"font-size:{size:.1f}px;color:#{o.get('color', INK)};font-weight:{700 if o.get('bold') else 400};"
        style += f"margin:0 0 {size*0.55:.1f}px 0;"
        if align:
            style += f"text-align:{align};"
        body = html.escape(text).replace("\n", "<br>")
        cls = ' class="bul"' if o.get("bullet") else ""
        out.append(f'<p{cls} style="{style}">{body}</p>')
    return "".join(out)


def build_html(path):
    pages = []
    for n, spec in enumerate(SLIDES, 1):
        parts = []
        if spec["title"]:
            parts.append(f'<div class="el" style="{css_box(0.7, 0.42, 11.9, 0.8)}">{html_paras([(spec["title"], {"size": 32, "bold": True})])}</div>')
            parts.append(f'<div class="el" style="{css_box(0.72, 1.18, 1.2, 0.06)}background:#{ACCENT}"></div>')
            parts.append(f'<div class="el" style="{css_box(0.7, 7.0, 6, 0.35)}">{html_paras([("VolterraDSPPredictor", {"size": 11, "color": MUTED})])}</div>')
            parts.append(f'<div class="el" style="{css_box(11.6, 7.0, 1.0, 0.35)}">{html_paras([(str(n), {"size": 11, "color": MUTED})], "center")}</div>')
        else:
            parts.append(f'<div class="el" style="{css_box(0, 0, 0.35, H)}background:#{ACCENT}"></div>')
        for el in spec["elements"]:
            kind = el[0]
            if kind == "text":
                _, x, y, w, h, paras, opts = el
                parts.append(f'<div class="el" style="{css_box(x, y, w, h)}">{html_paras(paras, opts.get("align"))}</div>')
            elif kind == "image":
                _, x, y, w, img, caption = el
                pw, ph = png_size(os.path.join(HERE, img))
                h = w * ph / pw
                parts.append(f'<img class="el" src="{img}" style="{css_box(x, y, w, h)}">')
                parts.append(f'<div class="el" style="{css_box(x, y + h + 0.08, w, 0.7)}">{html_paras([(caption, {"size": 13, "color": MUTED})])}</div>')
            elif kind == "tile":
                _, x, y, w, h, big, label, color, fill = el
                parts.append(f'<div class="el tile" style="{css_box(x, y, w, h)}background:#{fill};border-left:{0.09*PX:.1f}px solid #{color}"></div>')
                parts.append(f'<div class="el" style="{css_box(x + 0.3, y + 0.22, w - 0.5, h - 0.3)}">' +
                             html_paras([(big, {"size": 30 if len(big) < 14 else 26, "bold": True, "color": color}),
                                         (label, {"size": 15})]) + "</div>")
            elif kind == "box":
                _, x, y, w, h, title, sub, color, fill = el
                parts.append(f'<div class="el box" style="{css_box(x, y, w, h)}background:#{fill};border:2px solid #{color}">' +
                             html_paras([(title, {"size": 17, "bold": True}), (sub, {"size": 12, "color": MUTED})], "center") + "</div>")
            elif kind == "arrow":
                _, x1, y1, x2, y2, _label = el
                parts.append(f'<svg class="el" style="left:0;top:0;width:{W*PX}px;height:{H*PX}px;overflow:visible">'
                             f'<defs><marker id="m{n}" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto">'
                             f'<path d="M0,0 L8,4 L0,8 z" fill="#{INK}"/></marker></defs>'
                             f'<line x1="{x1*PX}" y1="{y1*PX}" x2="{x2*PX}" y2="{y2*PX}" stroke="#{INK}" stroke-width="2.6" marker-end="url(#m{n})"/></svg>')
        pages.append(f'<section class="slide">{"".join(parts)}</section>')

    doc = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>VolterraDSPPredictor — Slides</title>
<style>
@page {{ size: {W*PX:.0f}px {H*PX:.0f}px; margin: 0; }}
html, body {{ margin: 0; padding: 0; background: #888; }}
body {{ font-family: "{FONT}", -apple-system, Helvetica, Arial, sans-serif; color: #{INK}; }}
.slide {{ position: relative; width: {W*PX:.0f}px; height: {H*PX:.0f}px; background: white; overflow: hidden;
          break-after: page; margin: 0 auto 20px; }}
@media print {{ .slide {{ margin: 0; }} body {{ background: white; }} }}
.el {{ position: absolute; box-sizing: border-box; }}
.tile {{ border-radius: 8px; }}
.box {{ border-radius: 9px; display: flex; flex-direction: column; justify-content: center; padding: 4px; }}
.box p:last-child {{ margin-bottom: 0 !important; }}
img.el {{ border-radius: 6px; }}
p {{ line-height: 1.22; }}
p.bul {{ padding-left: 0.32in; text-indent: -0.26in; }}
p.bul::before {{ content: "•"; display: inline-block; width: 0.26in; text-indent: 0; }}
</style></head><body>{"".join(pages)}</body></html>"""
    with open(path, "w") as f:
        f.write(doc)


if __name__ == "__main__":
    build_pptx(os.path.join(HERE, "VolterraDSPPredictor_Slides.pptx"))
    build_html(os.path.join(HERE, "VolterraDSPPredictor_Slides.html"))
    print("wrote", len(SLIDES), "slides")
